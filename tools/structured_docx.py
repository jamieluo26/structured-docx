"""Finite editable text/heading/native-list DOCX generation and independent readback.

Only explicit bytes are accepted. Neither API reads caller files or accesses a network.
See docs/structured_docx.md for the supported source and package profile.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
import hashlib
import importlib.metadata
import json
import re
import struct
import sys
import zipfile
from xml.etree import ElementTree as ET
from xml.sax.saxutils import quoteattr


VERSION = "structured-docx/v1"
MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
HARD_LIMITS = {
    "source_bytes": 2 * 1024 * 1024, "text_chars": 500000,
    "string_chars": 32768, "blocks": 1000, "list_items": 2000,
    "list_depth": 4, "model_depth": 16, "archive_bytes": 8 * 1024 * 1024,
    "zip_members": 128, "uncompressed_bytes": 32 * 1024 * 1024,
    "xml_bytes": 4 * 1024 * 1024, "compression_ratio": 1000,
}
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/package/2006/relationships"
C = "http://schemas.openxmlformats.org/package/2006/content-types"
CP = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
DC = "http://purl.org/dc/elements/1.1/"
DT = "http://purl.org/dc/terms/"
XSI = "http://www.w3.org/2001/XMLSchema-instance"
APP = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
XML = "http://www.w3.org/XML/1998/namespace"
OFFICE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships/"
BINDING_PART = "docProps/structured-docx-binding.json"
BINDING_REL = "urn:structured-docx:relationships:binding"
PART_TYPES = {
    "word/document.xml": "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    "word/styles.xml": "application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml",
    "word/numbering.xml": "application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml",
    "docProps/core.xml": "application/vnd.openxmlformats-package.core-properties+xml",
    "docProps/app.xml": "application/vnd.openxmlformats-officedocument.extended-properties+xml",
    BINDING_PART: "application/json",
}
BASE_PARTS = frozenset(PART_TYPES) | {"[Content_Types].xml", "_rels/.rels", "word/_rels/document.xml.rels"}
THUMBNAIL = "docProps/thumbnail.jpeg"
PROFILE = {
    "version": VERSION, "page": "Letter portrait; one inch margins",
    "body_points": 11, "ascii_font": "Arial", "east_asia_font": "Microsoft YaHei",
    "title_points": 22, "heading_points": [16, 14, 12], "color": "000000",
    "whitespace": "CRLF and CR project to LF; all remaining characters preserved",
    "list": "four native levels; each list instance owns a new numId and restarts at one; mixed child kinds supported",
    "bullet_glyphs": ["•", "◦", "▪", "•"],
    "metadata_default": {"author": "", "language": "en-US", "createdAt": "2000-01-01T00:00:00Z"},
    "package": "finite nine-part package; optional bounded inert JPEG thumbnail accepted on verification",
    "opc_namespaces": "content-type and relationship roots use local default-namespace serialization",
    "binding": "raw source SHA256, canonical model SHA256, profile SHA256; consistency only, not a signature",
}


class DocumentError(ValueError):
    """Stable error without caller text or filesystem paths."""
    def __init__(self, code: str, path: str = "$", stage: str = "input"):
        self.code, self.path, self.stage = code, path, stage
        super().__init__(f"{stage}:{code}:{path}")


@dataclass(frozen=True)
class RenderResult:
    docx_bytes: bytes
    mime_type: str
    filename: str
    source_sha256: str
    model_sha256: str
    render_profile_sha256: str
    docx_sha256: str
    docx_size: int
    receipt: dict


@dataclass(frozen=True)
class Verification:
    status: str
    code: str
    source_sha256: str | None
    model_sha256: str | None
    render_profile_sha256: str
    docx_sha256: str | None
    docx_size: int
    mismatches: list
    unknowns: list
    observed: list
    receipt: dict


def _sha(data):
    return hashlib.sha256(data).hexdigest()


def _canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


RENDER_PROFILE_SHA256 = _sha(_canonical(PROFILE))


def _limits(overrides):
    answer = dict(HARD_LIMITS)
    if overrides is None:
        return answer
    if type(overrides) is not dict:
        raise DocumentError("LIMITS_TYPE")
    for key, value in list(overrides.items()):
        if key not in answer or type(value) is not int or not 1 <= value <= answer[key]:
            raise DocumentError("LIMITS_INVALID")
        answer[key] = value
    return answer


def _capture(value, cap, stage):
    if not isinstance(value, (bytes, bytearray, memoryview)):
        raise DocumentError("BYTES_REQUIRED", "$", stage)
    length = value.nbytes if isinstance(value, memoryview) else len(value)
    if length > cap:
        raise DocumentError("RESOURCE_LIMIT", "$", stage)
    try:
        return bytes(value)
    except Exception:
        raise DocumentError("BYTES_CAPTURE_FAILED", "$", stage) from None


def _hash_identity(value, name, allow_none=True):
    if value is None and allow_none:
        return
    if type(value) is not str or not re.fullmatch(r"[0-9a-f]{64}", value):
        raise DocumentError("IDENTITY_FORMAT", name, "identity")


def _decode_json(data, depth_cap, stage):
    try:
        source = data.decode("utf-8")
    except UnicodeDecodeError:
        raise DocumentError("UTF8_INVALID", "$", stage) from None
    depth = 0
    quoted = escaped = False
    for char in source:
        if quoted:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                quoted = False
        elif char == '"':
            quoted = True
        elif char in "[{":
            depth += 1
            if depth > depth_cap:
                raise DocumentError("MODEL_DEPTH_LIMIT", "$", stage)
        elif char in "]}":
            depth -= 1
    def pairs(values):
        output = {}
        for key, value in values:
            if key in output:
                raise DocumentError("DUPLICATE_KEY", "$", stage)
            output[key] = value
        return output
    def constant(_):
        raise DocumentError("NONFINITE_NUMBER", "$", stage)
    try:
        return json.loads(source, object_pairs_hook=pairs, parse_constant=constant)
    except DocumentError:
        raise
    except (ValueError, RecursionError, OverflowError):
        raise DocumentError("JSON_INVALID", "$", stage) from None


def _valid_xml_text(value):
    return all(c in "\t\n\r" or 0x20 <= ord(c) <= 0xD7FF or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF for c in value)


def _parse_source(source_bytes, limits):
    raw = _capture(source_bytes, limits["source_bytes"], "input")
    model = _decode_json(raw, limits["model_depth"], "input")
    totals = {"text_chars": 0, "list_items": 0, "lists": 0, "blocks": 0}
    def keys(obj, allowed, required, path):
        if type(obj) is not dict:
            raise DocumentError("OBJECT_REQUIRED", path)
        if set(obj) - set(allowed):
            raise DocumentError("UNKNOWN_FIELD", path)
        if set(required) - set(obj):
            raise DocumentError("MISSING_FIELD", path)
    def literal(value, path, cap=None):
        if type(value) is not str:
            raise DocumentError("TEXT_REQUIRED", path)
        if len(value) > (limits["string_chars"] if cap is None else min(cap, limits["string_chars"])):
            raise DocumentError("STRING_LIMIT", path)
        if not _valid_xml_text(value):
            raise DocumentError("XML_CHARACTER_INVALID", path)
        totals["text_chars"] += len(value)
        if totals["text_chars"] > limits["text_chars"]:
            raise DocumentError("TEXT_LIMIT", path)
        return value.replace("\r\n", "\n").replace("\r", "\n")
    def list_model(obj, path, level):
        keys(obj, ("type", "ordered", "items"), ("type", "ordered", "items"), path)
        if obj["type"] != "list" or type(obj["ordered"]) is not bool:
            raise DocumentError("LIST_TYPE", path)
        if level > limits["list_depth"]:
            raise DocumentError("LIST_DEPTH_LIMIT", path)
        if type(obj["items"]) is not list:
            raise DocumentError("LIST_ITEMS_TYPE", path + ".items")
        if not obj["items"]:
            raise DocumentError("LIST_EMPTY", path + ".items")
        totals["lists"] += 1
        result = {"type": "list", "ordered": obj["ordered"], "items": []}
        for index, item in enumerate(obj["items"]):
            ip = f"{path}.items[{index}]"
            keys(item, ("text", "children"), ("text",), ip)
            totals["list_items"] += 1
            if totals["list_items"] > limits["list_items"]:
                raise DocumentError("LIST_ITEM_LIMIT", ip)
            output = {"text": literal(item["text"], ip + ".text")}
            if "children" in item:
                output["children"] = list_model(item["children"], ip + ".children", level + 1)
            result["items"].append(output)
        return result
    keys(model, ("version", "title", "blocks", "author", "language", "createdAt"), ("version", "title", "blocks"), "$")
    if model["version"] != VERSION or type(model["version"]) is not str:
        raise DocumentError("VERSION_UNSUPPORTED", "$.version")
    if type(model["blocks"]) is not list:
        raise DocumentError("BLOCKS_TYPE", "$.blocks")
    if len(model["blocks"]) > limits["blocks"]:
        raise DocumentError("BLOCK_LIMIT", "$.blocks")
    normalized = {"version": VERSION, "title": literal(model["title"], "$.title"), "blocks": []}
    totals["blocks"] = len(model["blocks"])
    for index, block in enumerate(model["blocks"]):
        path = f"$.blocks[{index}]"
        if type(block) is not dict:
            raise DocumentError("OBJECT_REQUIRED", path)
        kind = block.get("type")
        if kind == "list":
            output = list_model(block, path, 1)
        elif kind == "heading":
            keys(block, ("type", "level", "text"), ("type", "level", "text"), path)
            if type(block["level"]) is not int or not 1 <= block["level"] <= 3:
                raise DocumentError("HEADING_LEVEL", path + ".level")
            output = {"type": "heading", "level": block["level"], "text": literal(block["text"], path + ".text")}
        elif kind == "paragraph":
            keys(block, ("type", "text"), ("type", "text"), path)
            output = {"type": "paragraph", "text": literal(block["text"], path + ".text")}
        else:
            raise DocumentError("BLOCK_UNSUPPORTED", path)
        normalized["blocks"].append(output)
    if "author" in model:
        normalized["author"] = literal(model["author"], "$.author", 255)
    if "language" in model:
        language = literal(model["language"], "$.language", 63)
        if not re.fullmatch(r"[A-Za-z]{2,8}(?:-[A-Za-z0-9]{1,8})*", language):
            raise DocumentError("LANGUAGE_INVALID", "$.language")
        normalized["language"] = language
    if "createdAt" in model:
        created = literal(model["createdAt"], "$.createdAt", 20)
        try:
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z", created):
                raise ValueError
            datetime.strptime(created, "%Y-%m-%dT%H:%M:%SZ")
        except ValueError:
            raise DocumentError("CREATED_AT_INVALID", "$.createdAt") from None
        normalized["createdAt"] = created
    return raw, normalized, totals


def _metadata(model):
    return {key: model.get(key, value) for key, value in PROFILE["metadata_default"].items()}


def _tag(local):
    return "{" + W + "}" + local


def _element(tag, attributes=None, text=None):
    e = ET.Element(tag, attributes or {})
    if text is not None:
        e.text = text
    return e


def _w(parent, name, attrs=None):
    e = ET.SubElement(parent, _tag(name))
    for key, value in (attrs or {}).items():
        e.set(_tag(key), str(value))
    return e


def _xml(root):
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _style_specs(language):
    normal_r = {"rFonts": {"ascii": "Arial", "hAnsi": "Arial", "eastAsia": "Microsoft YaHei", "cs": "Arial"}, "sz": {"val": "22"}, "color": {"val": "000000"}, "lang": {"val": language}}
    normal_p = {"spacing": {"before": "0", "after": "120", "line": "276", "lineRule": "auto"}, "jc": {"val": "left"}}
    return normal_r, normal_p


def _styles_part(language):
    root = _element(_tag("styles"))
    defaults = _w(root, "docDefaults")
    rp = _w(_w(defaults, "rPrDefault"), "rPr")
    pp = _w(_w(defaults, "pPrDefault"), "pPr")
    nr, np = _style_specs(language)
    for name, attrs in nr.items(): _w(rp, name, attrs)
    for name, attrs in np.items(): _w(pp, name, attrs)
    for sid, name, points, outline in [("Normal", "Normal", 11, None), ("Title", "Title", 22, None), ("Heading1", "heading 1", 16, 0), ("Heading2", "heading 2", 14, 1), ("Heading3", "heading 3", 12, 2), ("ListParagraph", "List Paragraph", 11, None)]:
        style = _w(root, "style", {"type": "paragraph", "styleId": sid, **({"default": "1"} if sid == "Normal" else {})})
        _w(style, "name", {"val": name})
        if sid != "Normal": _w(style, "basedOn", {"val": "Normal"})
        _w(style, "next", {"val": "Normal"})
        if sid in ("Title", "Heading1", "Heading2", "Heading3"): _w(style, "qFormat")
        p = _w(style, "pPr")
        for tag, attrs in np.items(): _w(p, tag, attrs)
        if sid in ("Title", "Heading1", "Heading2", "Heading3"): _w(p, "keepNext")
        if outline is not None: _w(p, "outlineLvl", {"val": outline})
        r = _w(style, "rPr")
        for tag, attrs in nr.items(): _w(r, tag, attrs)
        r.find(_tag("sz")).set(_tag("val"), str(points * 2))
        if sid in ("Title", "Heading1", "Heading2", "Heading3"): _w(r, "b")
    return _xml(root)


def _numbering_part(number_instances):
    root = _element(_tag("numbering"))
    for aid, ordered in ((1, True), (2, False)):
        abstract = _w(root, "abstractNum", {"abstractNumId": aid})
        _w(abstract, "multiLevelType", {"val": "multilevel"})
        for level in range(4):
            e = _w(abstract, "lvl", {"ilvl": level})
            _w(e, "start", {"val": 1})
            _w(e, "numFmt", {"val": "decimal" if ordered else "bullet"})
            _w(e, "lvlText", {"val": f"%{level + 1}." if ordered else PROFILE["bullet_glyphs"][level]})
            _w(e, "lvlJc", {"val": "left"})
            pp = _w(e, "pPr")
            tabs = _w(pp, "tabs")
            _w(tabs, "tab", {"val": "num", "pos": 360 * (level + 1)})
            _w(pp, "ind", {"left": 360 * (level + 1), "hanging": 180})
    for num_id, ordered in number_instances:
        num = _w(root, "num", {"numId": num_id})
        _w(num, "abstractNumId", {"val": 1 if ordered else 2})
        for level in range(4):
            _w(_w(num, "lvlOverride", {"ilvl": level}), "startOverride", {"val": 1})
    return _xml(root)


def _core_part(metadata):
    root = _element("{" + CP + "}coreProperties")
    for tag, value in [("{" + DC + "}creator", metadata["author"]), ("{" + DC + "}title", ""), ("{" + DC + "}subject", ""), ("{" + DC + "}description", ""), ("{" + DC + "}language", metadata["language"]), ("{" + CP + "}keywords", ""), ("{" + CP + "}lastModifiedBy", ""), ("{" + CP + "}revision", "1"), ("{" + CP + "}category", "")]:
        ET.SubElement(root, tag).text = value
    for local in ("created", "modified"):
        ET.SubElement(root, "{" + DT + "}" + local, {"{" + XSI + "}type": "dcterms:W3CDTF"}).text = metadata["createdAt"]
    # xsi:type uses the declared dcterms prefix, which has no element QName use.
    root.set("xmlns:dcterms", DT)
    return _xml(root)


def _opc_part(root_name, namespace, rows):
    """Serialize only fixed OPC tags locally; never mutate ET's global registry.

    OPC attributes are unqualified. ElementTree's default_namespace argument
    rejects those attributes, while a prefix-only OPC root is not accepted by
    some document renderers. All tag names here are fixed internal literals;
    attribute values are XML-quoted even though these parts contain no text.
    """
    allowed = {("Relationships", R): {"Relationship"}, ("Types", C): {"Default", "Override"}}
    tags = allowed.get((root_name, namespace))
    if tags is None or any(name not in tags for name, _ in rows):
        raise DocumentError("OPC_PROFILE_INVALID", "$", "generation")
    body = []
    for name, attributes in rows:
        attributes_text = "".join(" " + key + "=" + quoteattr(str(value)) for key, value in attributes.items())
        body.append("<" + name + attributes_text + "/>")
    return ("<?xml version=\"1.0\" encoding=\"UTF-8\"?>\n<" + root_name + " xmlns=" + quoteattr(namespace) + ">" + "".join(body) + "</" + root_name + ">").encode("utf-8")


def _relationships(rows):
    children = [("Relationship", {"Id": f"rId{index}", "Type": kind, "Target": target}) for index, (kind, target) in enumerate(rows, 1)]
    return _opc_part("Relationships", R, children)


def _content_types():
    children = [
        ("Default", {"Extension": "rels", "ContentType": "application/vnd.openxmlformats-package.relationships+xml"}),
        ("Default", {"Extension": "xml", "ContentType": "application/xml"}),
    ]
    for part, kind in PART_TYPES.items():
        children.append(("Override", {"PartName": "/" + part, "ContentType": kind}))
    return _opc_part("Types", C, children)


def _runtime_provenance():
    try:
        version = importlib.metadata.version("python-docx")
    except Exception:
        version = None
    return {"python": sys.version.split()[0], "python_docx": version, "verifier": "stdlib zipfile and ElementTree; independent body walk"}


def render_docx(source_bytes, *, expected_source_sha256=None, limits=None):
    """Generate one finite editable DOCX; never reads/writes caller paths."""
    cap = _limits(limits)
    raw, model, totals = _parse_source(source_bytes, cap)
    raw_sha, model_sha = _sha(raw), _sha(_canonical(model))
    _hash_identity(expected_source_sha256, "$expected_source_sha256")
    if expected_source_sha256 is not None and expected_source_sha256 != raw_sha:
        raise DocumentError("SOURCE_IDENTITY_MISMATCH", "$", "identity")
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError:
        raise DocumentError("DEPENDENCY_UNAVAILABLE", "$", "generation") from None
    number_instances = []
    try:
        document = Document()
        document.add_paragraph(model["title"], style="Title")
        def emit_list(value, level):
            if value["items"]:
                num_id = len(number_instances) + 1
                number_instances.append((num_id, value["ordered"]))
            else:
                return
            for item in value["items"]:
                p = document.add_paragraph(item["text"], style="List Paragraph")
                prop = OxmlElement("w:numPr")
                ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), str(level - 1)); prop.append(ilvl)
                num = OxmlElement("w:numId"); num.set(qn("w:val"), str(num_id)); prop.append(num)
                p._p.get_or_add_pPr().append(prop)
                if "children" in item:
                    emit_list(item["children"], level + 1)
        for block in model["blocks"]:
            if block["type"] == "heading":
                document.add_paragraph(block["text"], style=f"Heading {block['level']}")
            elif block["type"] == "paragraph":
                document.add_paragraph(block["text"], style="Normal")
            else:
                emit_list(block, 1)
        buffer = BytesIO()
        document.save(buffer)
        # Only the library's generated main-body XML is retained, not its template parts.
        with zipfile.ZipFile(BytesIO(buffer.getvalue())) as generated:
            body = ET.fromstring(generated.read("word/document.xml"))
        body.attrib.clear()
        main = body.find(_tag("body"))
        section = main.find(_tag("sectPr"))
        if section is None:
            section = ET.SubElement(main, _tag("sectPr"))
        section.clear()
        _w(section, "pgSz", {"w": 12240, "h": 15840})
        _w(section, "pgMar", {"top": 1440, "right": 1440, "bottom": 1440, "left": 1440, "header": 720, "footer": 720, "gutter": 0})
        metadata = _metadata(model)
        binding = {"version": VERSION, "source_sha256": raw_sha, "model_sha256": model_sha, "render_profile_sha256": RENDER_PROFILE_SHA256}
        app = _element("{" + APP + "}Properties")
        ET.SubElement(app, "{" + APP + "}Application").text = "structured-docx"
        parts = {
            "[Content_Types].xml": _content_types(),
            "_rels/.rels": _relationships([(OFFICE_REL + "officeDocument", "word/document.xml"), (PKG_REL + "metadata/core-properties", "docProps/core.xml"), (OFFICE_REL + "extended-properties", "docProps/app.xml"), (BINDING_REL, BINDING_PART)]),
            "word/document.xml": _xml(body),
            "word/_rels/document.xml.rels": _relationships([(OFFICE_REL + "styles", "styles.xml"), (OFFICE_REL + "numbering", "numbering.xml")]),
            "word/styles.xml": _styles_part(metadata["language"]),
            "word/numbering.xml": _numbering_part(number_instances),
            "docProps/core.xml": _core_part(metadata), "docProps/app.xml": _xml(app),
            BINDING_PART: _canonical(binding),
        }
        if len(parts) > cap["zip_members"] or sum(map(len, parts.values())) > cap["uncompressed_bytes"]:
            raise DocumentError("RESOURCE_LIMIT", "$", "generation")
        if any(len(value) > cap["xml_bytes"] for key, value in parts.items() if key.endswith((".xml", ".rels"))):
            raise DocumentError("XML_LIMIT", "$", "generation")
        output = BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as package:
            for name, value in parts.items():
                info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                package.writestr(info, value)
        actual = output.getvalue()
        if len(actual) > cap["archive_bytes"]:
            raise DocumentError("RESOURCE_LIMIT", "$", "generation")
        # Apply the same declared ZIP resource limits before returning success.
        _safe_package(actual, cap)
    except DocumentError:
        raise
    except Exception:
        raise DocumentError("GENERATION_FAILED", "$", "generation") from None
    output_sha = _sha(actual)
    receipt = {"version": VERSION, "stage": "generation", "status": "GENERATED_NOT_VERIFIED", "source_sha256": raw_sha, "source_size": len(raw), "model_sha256": model_sha, "render_profile_sha256": RENDER_PROFILE_SHA256, "docx_sha256": output_sha, "docx_size": len(actual), "supported_counts": totals, "dependencies": _runtime_provenance(), "semantic": "NOT_RUN", "byte_identity": "CAPTURED", "render": "NOT_RUN", "binding_authentication": "NONE"}
    return RenderResult(actual, MIME, "structured-document.docx", raw_sha, model_sha, RENDER_PROFILE_SHA256, output_sha, len(actual), receipt)


def _parse_xml(data, part, cap):
    if len(data) > cap["xml_bytes"]:
        raise DocumentError("XML_LIMIT", part, "verification")
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        raise DocumentError("XML_ENCODING", part, "verification") from None
    if re.search(r"<!\s*(?:DOCTYPE|ENTITY)", text, re.I) or "<!--" in text or re.search(r"<\?(?!xml\s)", text, re.I):
        raise DocumentError("XML_UNSUPPORTED_DECLARATION", part, "verification")
    match = re.match(r"\s*<\?xml\s+[^?]*encoding\s*=\s*['\"]([^'\"]+)", text, re.I)
    if match and match.group(1).lower() not in ("utf-8", "utf8"):
        raise DocumentError("XML_ENCODING", part, "verification")
    try:
        root = ET.fromstring(text)
    except (ET.ParseError, ValueError, RecursionError):
        raise DocumentError("XML_INVALID", part, "verification") from None
    stack, count = [(root, 1)], 0
    while stack:
        elem, depth = stack.pop(); count += 1
        if depth > 64 or count > 200000:
            raise DocumentError("XML_COMPLEXITY_LIMIT", part, "verification")
        stack.extend((child, depth + 1) for child in elem)
    return root


def _safe_package(data, cap):
    try:
        # This finite non-ZIP64 profile has no archive comment, prepended executable,
        # inter-member padding or appended payload. Do not let zipfile's permissive
        # handling of such bytes turn an uninspected tail into verified content.
        if len(data) < 22 or data[:4] != b"PK\x03\x04" or data[-22:-18] != b"PK\x05\x06":
            raise DocumentError("ZIP_ENVELOPE_UNSUPPORTED", "$docx", "verification")
        end = struct.unpack("<4s4H2LH", data[-22:])
        if end[1] or end[2] or end[3] != end[4] or end[7] or end[5] + end[6] != len(data) - 22:
            raise DocumentError("ZIP_ENVELOPE_UNSUPPORTED", "$docx", "verification")
        with zipfile.ZipFile(BytesIO(data)) as package:
            info = package.infolist()
            if len(info) != end[4] or package.comment:
                raise DocumentError("ZIP_ENVELOPE_UNSUPPORTED", "$docx", "verification")
            if len(info) > cap["zip_members"]:
                raise DocumentError("ZIP_MEMBER_LIMIT", "$docx", "verification")
            names, total, local_end = set(), 0, 0
            for item in info:
                name = item.filename
                if name in names:
                    raise DocumentError("ZIP_DUPLICATE_MEMBER", "$docx", "verification")
                names.add(name)
                if not re.fullmatch(r"[A-Za-z0-9_\[\]./-]+", name) or name.startswith("/") or any(x in ("", ".", "..") for x in name.split("/")):
                    raise DocumentError("ZIP_UNSAFE_PATH", "$docx", "verification")
                if item.flag_bits & ~0x800 or item.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                    raise DocumentError("ZIP_UNSUPPORTED_COMPRESSION", "$docx", "verification")
                if item.extra or item.comment or item.header_offset != local_end or local_end + 30 > len(data):
                    raise DocumentError("ZIP_ENVELOPE_UNSUPPORTED", "$docx", "verification")
                header = struct.unpack("<4s5H3L2H", data[local_end:local_end + 30])
                if header[0] != b"PK\x03\x04" or header[2] != item.flag_bits or header[3] != item.compress_type or header[6] != item.CRC or header[7] != item.compress_size or header[8] != item.file_size or header[10]:
                    raise DocumentError("ZIP_ENVELOPE_UNSUPPORTED", "$docx", "verification")
                name_bytes = data[local_end + 30:local_end + 30 + header[9]]
                if name_bytes != item.filename.encode("utf-8"):
                    raise DocumentError("ZIP_UNSAFE_PATH", "$docx", "verification")
                local_end += 30 + header[9] + header[10] + item.compress_size
                if item.file_size < 0 or item.file_size > cap["uncompressed_bytes"] or item.compress_size < 0:
                    raise DocumentError("ZIP_SIZE_LIMIT", "$docx", "verification")
                if item.file_size and (not item.compress_size or item.file_size > max(1, item.compress_size) * cap["compression_ratio"]):
                    raise DocumentError("ZIP_COMPRESSION_RATIO", "$docx", "verification")
                if (item.external_attr >> 16) & 0o170000 in (0o120000, 0o040000):
                    raise DocumentError("ZIP_UNSUPPORTED_MEMBER_TYPE", "$docx", "verification")
                total += item.file_size
                if total > cap["uncompressed_bytes"]:
                    raise DocumentError("ZIP_SIZE_LIMIT", "$docx", "verification")
                if name.endswith((".xml", ".rels")) and item.file_size > cap["xml_bytes"]:
                    raise DocumentError("XML_LIMIT", "$docx", "verification")
                if name == BINDING_PART and item.file_size > 4096:
                    raise DocumentError("BINDING_LIMIT", "$docx", "verification")
            if local_end != end[6]:
                raise DocumentError("ZIP_ENVELOPE_UNSUPPORTED", "$docx", "verification")
            if names - BASE_PARTS - {THUMBNAIL}:
                raise DocumentError("PACKAGE_PART_UNSUPPORTED", "$docx", "verification")
            if BASE_PARTS - names:
                raise DocumentError("PACKAGE_PART_MISSING", "$docx", "verification")
            result = {}
            for item in info:
                with package.open(item) as entry:
                    body = entry.read(item.file_size + 1)
                if len(body) != item.file_size:
                    raise DocumentError("ZIP_SIZE_MISMATCH", "$docx", "verification")
                result[item.filename] = body
            if THUMBNAIL in result:
                image = result[THUMBNAIL]
                if len(image) > 65536 or not image.startswith(b"\xff\xd8\xff") or not image.endswith(b"\xff\xd9"):
                    raise DocumentError("THUMBNAIL_UNSUPPORTED", "$docx", "verification")
            return result
    except DocumentError:
        raise
    except (zipfile.BadZipFile, RuntimeError, EOFError, OSError, ValueError, NotImplementedError):
        raise DocumentError("ZIP_INVALID", "$docx", "verification") from None


def _container(e, tag, attrs, part):
    if e.tag != tag or e.attrib != attrs or (e.text or "").strip() or any((c.tail or "").strip() for c in e):
        raise DocumentError("STRUCTURE_UNSUPPORTED", part, "verification")


def _leaf(e, tag, attrs, part):
    if e.tag != tag or e.attrib != attrs or len(e) or (e.text or "").strip():
        raise DocumentError("STRUCTURE_UNSUPPORTED", part, "verification")


def _single(e, name, part):
    found = e.findall(name)
    if len(found) != 1:
        raise DocumentError("STRUCTURE_MISSING_OR_DUPLICATE", part, "verification")
    return found[0]


def _read_relationships(root, expected, part):
    _container(root, "{" + R + "}Relationships", {}, part)
    observed, ids = set(), set()
    for child in root:
        if child.tag != "{" + R + "}Relationship" or set(child.attrib) != {"Id", "Type", "Target"} or len(child) or (child.text or "").strip():
            raise DocumentError("RELATIONSHIP_UNSUPPORTED", part, "verification")
        if not child.get("Id") or child.get("Id") in ids:
            raise DocumentError("RELATIONSHIP_DUPLICATE", part, "verification")
        ids.add(child.get("Id"))
        pair = (child.get("Type"), child.get("Target"))
        if pair in observed:
            raise DocumentError("RELATIONSHIP_DUPLICATE", part, "verification")
        observed.add(pair)
    if observed != set(expected):
        raise DocumentError("RELATIONSHIP_UNSUPPORTED", part, "verification")


def _read_package_structure(parts, cap):
    parsed = {name: _parse_xml(body, name, cap) for name, body in parts.items() if name.endswith((".xml", ".rels"))}
    ctypes = parsed["[Content_Types].xml"]
    _container(ctypes, "{" + C + "}Types", {}, "[Content_Types].xml")
    defaults, overrides = {}, {}
    for child in ctypes:
        if child.tag == "{" + C + "}Default" and set(child.attrib) == {"Extension", "ContentType"}:
            target, key, value = defaults, child.get("Extension"), child.get("ContentType")
        elif child.tag == "{" + C + "}Override" and set(child.attrib) == {"PartName", "ContentType"}:
            target, key, value = overrides, child.get("PartName"), child.get("ContentType")
        else:
            raise DocumentError("CONTENT_TYPE_UNSUPPORTED", "[Content_Types].xml", "verification")
        if key in target or len(child) or (child.text or "").strip():
            raise DocumentError("CONTENT_TYPE_DUPLICATE", "[Content_Types].xml", "verification")
        target[key] = value
    wanted_defaults = {"rels": "application/vnd.openxmlformats-package.relationships+xml", "xml": "application/xml"}
    if THUMBNAIL in parts:
        wanted_defaults["jpeg"] = "image/jpeg"
    if defaults != wanted_defaults or overrides != {"/" + k: v for k, v in PART_TYPES.items()}:
        raise DocumentError("CONTENT_TYPE_UNSUPPORTED", "[Content_Types].xml", "verification")
    root_rels = [(OFFICE_REL + "officeDocument", "word/document.xml"), (PKG_REL + "metadata/core-properties", "docProps/core.xml"), (OFFICE_REL + "extended-properties", "docProps/app.xml"), (BINDING_REL, BINDING_PART)]
    if THUMBNAIL in parts:
        root_rels.append((PKG_REL + "metadata/thumbnail", THUMBNAIL))
    _read_relationships(parsed["_rels/.rels"], root_rels, "_rels/.rels")
    _read_relationships(parsed["word/_rels/document.xml.rels"], [(OFFICE_REL + "styles", "styles.xml"), (OFFICE_REL + "numbering", "numbering.xml")], "word/_rels/document.xml.rels")
    return parsed


def _expected_inventory(model):
    """Independent iterative source walk: not used by the renderer."""
    result = [{"path": "$.title", "kind": "title", "level": 0, "text": model["title"]}]
    pending = [("block", b, f"$.blocks[{i}]", 0, None, None) for i, b in reversed(list(enumerate(model["blocks"])))]
    while pending:
        kind, value, path, level, list_id, index = pending.pop()
        if kind == "item":
            result.append({"path": path, "kind": "list", "level": level, "ordered": value["ordered"], "list_id": list_id, "item_index": index, "text": value["item"]["text"]})
            if "children" in value["item"]:
                pending.append(("block", value["item"]["children"], path + ".children", level, None, None))
        elif value["type"] == "list":
            for i in range(len(value["items"]) - 1, -1, -1):
                pending.append(("item", {"item": value["items"][i], "ordered": value["ordered"]}, f"{path}.items[{i}]", level + 1, path, i))
        else:
            result.append({"path": path, "kind": value["type"], "level": value.get("level", 0), "text": value["text"]})
    return result


def _read_numbering(root):
    part = "word/numbering.xml"
    _container(root, _tag("numbering"), {}, part)
    abstracts, nums = {}, {}
    for child in root:
        if child.tag == _tag("abstractNum"):
            aid = child.get(_tag("abstractNumId"))
            if child.attrib != {_tag("abstractNumId"): aid} or aid not in ("1", "2") or aid in abstracts:
                raise DocumentError("NUMBERING_UNSUPPORTED", part, "verification")
            abstract = {}
            if len(child) != 5:
                raise DocumentError("NUMBERING_UNSUPPORTED", part, "verification")
            _leaf(_single(child, _tag("multiLevelType"), part), _tag("multiLevelType"), {_tag("val"): "multilevel"}, part)
            for lvl in child.findall(_tag("lvl")):
                level = lvl.get(_tag("ilvl"))
                if level not in ("0", "1", "2", "3") or level in abstract or lvl.attrib != {_tag("ilvl"): level} or len(lvl) != 5:
                    raise DocumentError("NUMBERING_LEVEL", part, "verification")
                i = int(level)
                for local, attrs in [("start", {"val": "1"}), ("numFmt", {"val": "decimal" if aid == "1" else "bullet"}), ("lvlText", {"val": f"%{i + 1}." if aid == "1" else ["•", "◦", "▪", "•"][i]}), ("lvlJc", {"val": "left"})]:
                    _leaf(_single(lvl, _tag(local), part), _tag(local), {_tag(k): v for k, v in attrs.items()}, part)
                pp = _single(lvl, _tag("pPr"), part)
                _container(pp, _tag("pPr"), {}, part)
                if len(pp) != 2:
                    raise DocumentError("NUMBERING_LAYOUT", part, "verification")
                tabs = _single(pp, _tag("tabs"), part)
                _container(tabs, _tag("tabs"), {}, part)
                if len(tabs) != 1:
                    raise DocumentError("NUMBERING_LAYOUT", part, "verification")
                _leaf(tabs[0], _tag("tab"), {_tag("val"): "num", _tag("pos"): str(360 * (i + 1))}, part)
                _leaf(_single(pp, _tag("ind"), part), _tag("ind"), {_tag("left"): str(360 * (i + 1)), _tag("hanging"): "180"}, part)
                abstract[level] = {"ordered": aid == "1", "start": 1}
            if set(abstract) != {"0", "1", "2", "3"}:
                raise DocumentError("NUMBERING_LEVEL", part, "verification")
            abstracts[aid] = abstract
        elif child.tag == _tag("num"):
            num_id = child.get(_tag("numId"))
            if not num_id or not re.fullmatch(r"[1-9][0-9]*", num_id) or num_id in nums or child.attrib != {_tag("numId"): num_id} or len(child) != 5:
                raise DocumentError("NUMBERING_ID", part, "verification")
            aid_node = _single(child, _tag("abstractNumId"), part)
            aid = aid_node.get(_tag("val"))
            _leaf(aid_node, _tag("abstractNumId"), {_tag("val"): aid}, part)
            if aid not in ("1", "2"):
                raise DocumentError("NUMBERING_REFERENCE", part, "verification")
            levels = set()
            for override in child.findall(_tag("lvlOverride")):
                level = override.get(_tag("ilvl"))
                if level not in ("0", "1", "2", "3") or level in levels or override.attrib != {_tag("ilvl"): level} or len(override) != 1:
                    raise DocumentError("NUMBERING_RESTART", part, "verification")
                _leaf(override[0], _tag("startOverride"), {_tag("val"): "1"}, part)
                levels.add(level)
            if levels != {"0", "1", "2", "3"}:
                raise DocumentError("NUMBERING_RESTART", part, "verification")
            nums[num_id] = aid
        else:
            raise DocumentError("NUMBERING_UNSUPPORTED", part, "verification")
    if set(abstracts) != {"1", "2"}:
        raise DocumentError("NUMBERING_REFERENCE", part, "verification")
    return nums


def _property_map(node, part):
    if node.attrib or (node.text or "").strip():
        raise DocumentError("STYLE_UNSUPPORTED", part, "verification")
    result = {}
    for child in node:
        if not child.tag.startswith("{" + W + "}") or len(child) or (child.text or "").strip() or child.tag in result:
            raise DocumentError("STYLE_UNSUPPORTED", part, "verification")
        result[child.tag] = dict(child.attrib)
    return result


def _read_styles(root, language):
    part = "word/styles.xml"
    _container(root, _tag("styles"), {}, part)
    defaults = _single(root, _tag("docDefaults"), part)
    _container(defaults, _tag("docDefaults"), {}, part)
    if len(defaults) != 2:
        raise DocumentError("STYLE_UNSUPPORTED", part, "verification")
    # Explicit verifier profile, independent of the style XML generator.
    rp = {_tag("rFonts"): {_tag("ascii"): "Arial", _tag("hAnsi"): "Arial", _tag("eastAsia"): "Microsoft YaHei", _tag("cs"): "Arial"}, _tag("sz"): {_tag("val"): "22"}, _tag("color"): {_tag("val"): "000000"}, _tag("lang"): {_tag("val"): language}}
    pp = {_tag("spacing"): {_tag("before"): "0", _tag("after"): "120", _tag("line"): "276", _tag("lineRule"): "auto"}, _tag("jc"): {_tag("val"): "left"}}
    for parent, local, expected in (("rPrDefault", "rPr", rp), ("pPrDefault", "pPr", pp)):
        container = _single(defaults, _tag(parent), part)
        _container(container, _tag(parent), {}, part)
        if len(container) != 1 or _property_map(_single(container, _tag(local), part), part) != expected:
            raise DocumentError("STYLE_PROFILE_MISMATCH", part, "verification")
    specs = {"Normal": ("Normal", 22, None), "Title": ("Title", 44, None), "Heading1": ("heading 1", 32, "0"), "Heading2": ("heading 2", 28, "1"), "Heading3": ("heading 3", 24, "2"), "ListParagraph": ("List Paragraph", 22, None)}
    seen = set()
    for style in root:
        if style is defaults:
            continue
        sid = style.get(_tag("styleId"))
        if sid not in specs or sid in seen:
            raise DocumentError("STYLE_UNSUPPORTED", part, "verification")
        seen.add(sid)
        attrs = {_tag("type"): "paragraph", _tag("styleId"): sid}
        if sid == "Normal": attrs[_tag("default")] = "1"
        _container(style, _tag("style"), attrs, part)
        name, size, outline = specs[sid]
        special = sid in ("Title", "Heading1", "Heading2", "Heading3")
        if len(style) != 4 + (sid != "Normal") + special:
            raise DocumentError("STYLE_UNSUPPORTED", part, "verification")
        _leaf(_single(style, _tag("name"), part), _tag("name"), {_tag("val"): name}, part)
        _leaf(_single(style, _tag("next"), part), _tag("next"), {_tag("val"): "Normal"}, part)
        if sid != "Normal": _leaf(_single(style, _tag("basedOn"), part), _tag("basedOn"), {_tag("val"): "Normal"}, part)
        if special: _leaf(_single(style, _tag("qFormat"), part), _tag("qFormat"), {}, part)
        expected_p = {k: dict(v) for k, v in pp.items()}
        if special: expected_p[_tag("keepNext")] = {}
        if outline is not None: expected_p[_tag("outlineLvl")] = {_tag("val"): outline}
        expected_r = {k: dict(v) for k, v in rp.items()}
        expected_r[_tag("sz")] = {_tag("val"): str(size)}
        if special: expected_r[_tag("b")] = {}
        if _property_map(_single(style, _tag("pPr"), part), part) != expected_p or _property_map(_single(style, _tag("rPr"), part), part) != expected_r:
            raise DocumentError("STYLE_PROFILE_MISMATCH", part, "verification")
    if seen != set(specs):
        raise DocumentError("STYLE_MISSING", part, "verification")


def _read_metadata(parsed, expected):
    part = "docProps/core.xml"
    root = parsed[part]
    _container(root, "{" + CP + "}coreProperties", {}, part)
    wanted = {"{" + DC + "}creator": (expected["author"], {}), "{" + DC + "}title": ("", {}), "{" + DC + "}subject": ("", {}), "{" + DC + "}description": ("", {}), "{" + DC + "}language": (expected["language"], {}), "{" + CP + "}keywords": ("", {}), "{" + CP + "}lastModifiedBy": ("", {}), "{" + CP + "}revision": ("1", {}), "{" + CP + "}category": ("", {}), "{" + DT + "}created": (expected["createdAt"], {"{" + XSI + "}type": "dcterms:W3CDTF"}), "{" + DT + "}modified": (expected["createdAt"], {"{" + XSI + "}type": "dcterms:W3CDTF"})}
    seen = set()
    for item in root:
        if item.tag not in wanted or item.tag in seen or len(item):
            raise DocumentError("METADATA_UNSUPPORTED", part, "verification")
        value, attrs = wanted[item.tag]
        if item.attrib != attrs or (item.text or "") != value:
            raise DocumentError("METADATA_MISMATCH", part, "verification")
        seen.add(item.tag)
    if seen != set(wanted):
        raise DocumentError("METADATA_MISSING", part, "verification")
    app = parsed["docProps/app.xml"]
    _container(app, "{" + APP + "}Properties", {}, "docProps/app.xml")
    if len(app) != 1 or app[0].tag != "{" + APP + "}Application" or app[0].attrib or len(app[0]) or app[0].text != "structured-docx":
        raise DocumentError("METADATA_UNSUPPORTED", "docProps/app.xml", "verification")


def _read_body(root, numbering):
    part = "word/document.xml"
    _container(root, _tag("document"), {}, part)
    body = _single(root, _tag("body"), part)
    if len(root) != 1:
        raise DocumentError("BODY_UNSUPPORTED", part, "verification")
    _container(body, _tag("body"), {}, part)
    if not len(body) or body[-1].tag != _tag("sectPr"):
        raise DocumentError("SECTION_MISSING", part, "verification")
    section = body[-1]
    _container(section, _tag("sectPr"), {}, part)
    if len(section) != 2:
        raise DocumentError("SECTION_UNSUPPORTED", part, "verification")
    _leaf(_single(section, _tag("pgSz"), part), _tag("pgSz"), {_tag("w"): "12240", _tag("h"): "15840"}, part)
    _leaf(_single(section, _tag("pgMar"), part), _tag("pgMar"), {_tag(k): v for k, v in {"top": "1440", "right": "1440", "bottom": "1440", "left": "1440", "header": "720", "footer": "720", "gutter": "0"}.items()}, part)
    actual, counters = [], {}
    for index, paragraph in enumerate(body[:-1]):
        path = f"$docx.body[{index}]"
        _container(paragraph, _tag("p"), {}, path)
        props = paragraph.findall(_tag("pPr"))
        if len(props) > 1:
            raise DocumentError("PARAGRAPH_PROPERTIES", path, "verification")
        sid, num_id, level = "Normal", None, 0
        if props:
            prop = props[0]
            _container(prop, _tag("pPr"), {}, path)
            seen = set()
            for item in prop:
                if item.tag in seen:
                    raise DocumentError("PARAGRAPH_PROPERTIES", path, "verification")
                seen.add(item.tag)
                if item.tag == _tag("pStyle"):
                    sid = item.get(_tag("val"))
                    _leaf(item, _tag("pStyle"), {_tag("val"): sid}, path)
                elif item.tag == _tag("numPr"):
                    _container(item, _tag("numPr"), {}, path)
                    if len(item) != 2:
                        raise DocumentError("LIST_PROPERTIES", path, "verification")
                    ilvl, num = _single(item, _tag("ilvl"), path), _single(item, _tag("numId"), path)
                    value, num_id = ilvl.get(_tag("val")), num.get(_tag("val"))
                    _leaf(ilvl, _tag("ilvl"), {_tag("val"): value}, path)
                    _leaf(num, _tag("numId"), {_tag("val"): num_id}, path)
                    if value not in ("0", "1", "2", "3") or num_id not in numbering:
                        raise DocumentError("LIST_REFERENCE", path, "verification")
                    level = int(value) + 1
                else:
                    raise DocumentError("PARAGRAPH_PROPERTIES", path, "verification")
        text_parts = []
        for run in paragraph:
            if run in props:
                continue
            _container(run, _tag("r"), {}, path)
            for item in run:
                if item.tag == _tag("t"):
                    value = item.text or ""
                    if len(item) or item.attrib not in ({}, {"{" + XML + "}space": "preserve"}):
                        raise DocumentError("RUN_TEXT_UNSUPPORTED", path, "verification")
                    if value.strip() != value and item.get("{" + XML + "}space") != "preserve":
                        raise DocumentError("SPACE_PRESERVATION_MISSING", path, "verification")
                    text_parts.append(value)
                elif item.tag in (_tag("tab"), _tag("br")):
                    _leaf(item, item.tag, {}, path)
                    text_parts.append("\t" if item.tag == _tag("tab") else "\n")
                elif item.tag == _tag("rPr") and not item.attrib and not len(item) and not (item.text or "").strip():
                    pass
                else:
                    raise DocumentError("RUN_CONTENT_UNSUPPORTED", path, "verification")
        if num_id is not None:
            if sid != "ListParagraph":
                raise DocumentError("LIST_STYLE", path, "verification")
            key = (num_id, level)
            counters[key] = counters.get(key, 0) + 1
            item = {"path": path, "kind": "list", "level": level, "ordered": numbering[num_id] == "1", "num_id": num_id, "ordinal": counters[key], "text": "".join(text_parts)}
        else:
            if sid == "Title": kind, level = "title", 0
            elif sid in ("Heading1", "Heading2", "Heading3"): kind, level = "heading", int(sid[-1])
            elif sid == "Normal": kind, level = "paragraph", 0
            else: raise DocumentError("PARAGRAPH_STYLE_UNSUPPORTED", path, "verification")
            item = {"path": path, "kind": kind, "level": level, "text": "".join(text_parts)}
        actual.append(item)
    if set(numbering) != {item["num_id"] for item in actual if item["kind"] == "list"}:
        raise DocumentError("NUMBERING_UNUSED_OR_MISSING", "word/numbering.xml", "verification")
    return actual


def verify_docx(source_bytes, docx_bytes, *, expected_docx_sha256, expected_render_profile_sha256, limits=None):
    """Independently inspect exact selected package and compare every supported item."""
    source_sha = model_sha = output_sha = None
    size, observed, mismatches, unknowns = 0, [], [], []
    counts = None
    def mismatch(code, path, **details):
        if len(mismatches) < 128:
            mismatches.append({"code": code, "path": path, **details})
        elif not any(x["code"] == "DIAGNOSTIC_LIMIT" for x in unknowns):
            unknowns.append({"code": "DIAGNOSTIC_LIMIT", "path": "$verification"})
    try:
        cap = _limits(limits)
        raw, model, counts = _parse_source(source_bytes, cap)
        source_sha, model_sha = _sha(raw), _sha(_canonical(model))
        actual = _capture(docx_bytes, cap["archive_bytes"], "verification")
        size, output_sha = len(actual), _sha(actual)
        _hash_identity(expected_docx_sha256, "$expected_docx_sha256")
        _hash_identity(expected_render_profile_sha256, "$expected_render_profile_sha256")
        if expected_docx_sha256 is None or expected_render_profile_sha256 is None:
            unknowns.append({"code": "MISSING_EXPECTED_IDENTITY", "path": "$identity"})
        if expected_docx_sha256 is not None and output_sha != expected_docx_sha256:
            mismatch("OUTPUT_IDENTITY_MISMATCH", "$docx")
        if expected_render_profile_sha256 is not None and expected_render_profile_sha256 != RENDER_PROFILE_SHA256:
            mismatch("PROFILE_IDENTITY_MISMATCH", "$profile")
        parts = _safe_package(actual, cap)
        parsed = _read_package_structure(parts, cap)
        binding = _decode_json(parts[BINDING_PART], 4, "verification")
        expected_binding = {"version": VERSION, "source_sha256": source_sha, "model_sha256": model_sha, "render_profile_sha256": RENDER_PROFILE_SHA256}
        if type(binding) is not dict or set(binding) != set(expected_binding) or any(type(binding[k]) is not str for k in expected_binding) or binding != expected_binding:
            mismatch("SOURCE_BINDING_MISMATCH", BINDING_PART)
        meta = _metadata(model)
        _read_metadata(parsed, meta)
        _read_styles(parsed["word/styles.xml"], meta["language"])
        numbering = _read_numbering(parsed["word/numbering.xml"])
        observed = _read_body(parsed["word/document.xml"], numbering)
        expected = _expected_inventory(model)
        if len(observed) != len(expected):
            mismatch("ITEM_COUNT_MISMATCH", "$docx.body", expected_count=len(expected), observed_count=len(observed))
        list_to_num, num_to_list = {}, {}
        for index in range(min(len(expected), len(observed))):
            wanted, found = expected[index], observed[index]
            if wanted["kind"] != found["kind"] or wanted["level"] != found["level"]:
                mismatch("ITEM_STRUCTURE_MISMATCH", wanted["path"])
            if wanted["text"] != found["text"]:
                mismatch("ITEM_TEXT_MISMATCH", wanted["path"], expected_text_sha256=_sha(wanted["text"].encode("utf-8")), observed_text_sha256=_sha(found["text"].encode("utf-8")))
            if wanted["kind"] == "list" and found["kind"] == "list":
                lid, nid = wanted["list_id"], found["num_id"]
                if wanted["ordered"] is not found["ordered"]:
                    mismatch("LIST_KIND_MISMATCH", wanted["path"])
                if lid in list_to_num and list_to_num[lid] != nid or nid in num_to_list and num_to_list[nid] != lid:
                    mismatch("LIST_INSTANCE_MISMATCH", wanted["path"])
                list_to_num[lid], num_to_list[nid] = nid, lid
                if found["ordinal"] != wanted["item_index"] + 1:
                    mismatch("LIST_ORDER_OR_RESTART_MISMATCH", wanted["path"])
    except DocumentError as error:
        entry = {"code": error.code, "path": error.path, "stage": error.stage}
        if error.stage in ("input", "identity"):
            mismatch(error.code, error.path)
        else:
            unknowns.append(entry)
    except Exception:
        unknowns.append({"code": "VERIFICATION_UNAVAILABLE", "path": "$verification"})
    status = "FAIL" if mismatches else "NOT_VERIFIED" if unknowns else "PASS"
    code = mismatches[0]["code"] if mismatches else unknowns[0]["code"] if unknowns else "VERIFIED"
    receipt = {"version": VERSION, "stage": "verification", "status": status, "code": code, "source_sha256": source_sha, "model_sha256": model_sha, "render_profile_sha256": RENDER_PROFILE_SHA256, "docx_sha256": output_sha, "docx_size": size, "supported_counts": counts, "observed_items": len(observed), "mismatches": mismatches, "unknowns": unknowns, "dependencies": _runtime_provenance(), "semantic": "PASS" if status == "PASS" else "NOT_VERIFIED" if unknowns and not mismatches else "FAIL", "byte_identity": "PASS" if expected_docx_sha256 is not None and output_sha == expected_docx_sha256 else "NOT_VERIFIED" if expected_docx_sha256 is None else "FAIL", "render": "NOT_RUN", "binding_authentication": "NONE"}
    return Verification(status, code, source_sha, model_sha, RENDER_PROFILE_SHA256, output_sha, size, mismatches, unknowns, observed, receipt)
